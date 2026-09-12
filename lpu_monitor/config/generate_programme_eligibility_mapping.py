# """
# generate_programme_eligibility_mapping.py (SIMPLIFIED OUTPUT VERSION)

# Outputs a clean, 3-to-5 column mapping CSV (official_code, eligibility_base, row_label, status, reason)
# to perfectly match the expected input format for build_eligibility_expected.py.
# """
# import csv
# import glob
# import os
# import re
# import sys
# import pandas as pd
# import tempfile
# from collections import Counter
# from difflib import SequenceMatcher
# import openpyxl

# # Auto-discovered from incoming/ folders
# INCOMING_ELIGIBILITY_EXCEL = "lpu_monitor/incoming/eligibility_excel"
# INCOMING_PROGLIST = "lpu_monitor/incoming/programme_list"

# # Updated to match the expected naming convention
# OUTPUT_CSV = "lpu_monitor/config/eligibility_row_mapping.csv"

# EXACT_THRESHOLD = 0.90
# FUZZY_THRESHOLD = 0.55
# AMBIGUOUS_MARGIN = 0.06

# # Branch keyword groups for matching
# BRANCH_KEYWORDS = {
#     'cse': ['computer science', 'cse', 'ai and ml', 'artificial intelligence', 'machine learning',
#             'data science', 'data engineering', 'data analytics', 'cloud computing', 'generative ai',
#             'cyber security', 'block chain', 'robotics', 'software product', 'business systems',
#             'system design', 'full stack', 'user experience', 'ux', 'ui', 'big data', 'devops',
#             'information technology', 'it'],
#     'ece': ['electronics and communication', 'ece', 'semiconductor', 'vlsi', 'wireless', 'electronics'],
#     'ee': ['electrical engineering', 'ee', 'electrical'],
#     'eee': ['electrical and electronics', 'eee'],
#     'mech': ['mechanical', 'mechatronics', 'aerospace', 'automobile', 'me'],
#     'civil': ['civil', 'construction', 'ce'],
#     'chem': ['chemical'],
#     'bio': ['biotechnology', 'biomedical', 'food tech', 'genetics', 'bt', 'biotech'],
#     'mba': ['mba', 'business administration'],
#     'bba': ['bba', 'business administration'],
#     'bcom': ['b.com', 'commerce'],
#     'bca': ['bca', 'computer applications'],
#     'mca': ['mca', 'computer applications'],
#     'hotel': ['hotel', 'hospitality', 'catering', 'bhmct', 'tourism'],
#     'design': ['design', 'fashion', 'interior', 'animation', 'vfx', 'graphics', 'gaming', 'product', 'industrial'],
#     'law': ['law', 'llb', 'llm'],
#     'pharma': ['pharmacy', 'pharmaceutics', 'pharmacology', 'bpharm', 'mpharm', 'pharm d'],
# }

# # Parent category patterns (general eligibility rows that cover specializations)
# PARENT_PATTERNS = [
#     (r'b\.?\s?tech.*cse.*ece.*ee.*eee.*it', 'btech_cse_general'),
#     (r'b\.?\s?tech.*chemical.*civil.*mechanical', 'btech_mech_general'),
#     (r'b\.?\s?tech.*biotechnology.*food.*biomedical', 'btech_bio_general'),
#     (r'mba.*\(2yrs\)', 'mba_general'),
#     (r'mba.*\(2\s?yrs\)', 'mba_general'),
#     (r'bba.*\(3yrs\)', 'bba_general'),
#     (r'bba.*\(3\s?yrs\)', 'bba_general'),
#     (r'b\.?\s?com\.\s*$', 'bcom_general'),
#     (r'bachelor.*commerce(?!.*hons)', 'bcom_general'),
# ]

# # ---------------------------------------------------------------------------
# # .xls -> .xlsx conversion
# # ---------------------------------------------------------------------------
# def ensure_xlsx(path, workdir):
#     if path.lower().endswith(".xlsx"):
#         return path
#     if not path.lower().endswith(".xls"):
#         raise ValueError(f"Unsupported file type: {path}")

#     os.makedirs(workdir, exist_ok=True)
#     out_name = os.path.splitext(os.path.basename(path))[0] + ".xlsx"
#     out_path = os.path.join(workdir, out_name)

#     xls = pd.ExcelFile(path, engine="xlrd")
#     with pd.ExcelWriter(out_path, engine="openpyxl") as writer:
#         for sheet_name in xls.sheet_names:
#             df = xls.parse(sheet_name, header=None)
#             df.to_excel(writer, sheet_name=sheet_name, index=False, header=False)
#     return out_path

# # ---------------------------------------------------------------------------
# # Generic programme-name matching helpers
# # ---------------------------------------------------------------------------
# STOPWORDS = {
#     "b", "m", "of", "and", "the", "in", "for", "a", "an", "with", "hons",
#     "bachelor", "bachelors", "master", "masters", "programme", "program",
#     "technology", "engineering", "science", "sciences", "studies", "design",
#     "no", "integrated", "integr", "including", "option", "transfer", "credit",
#     "after", "exit", "discontinue", "lateral", "entry"
# }

# def _stem(t):
#     if len(t) > 4 and t.endswith("s") and not t.endswith("ss"):
#         return t[:-1]
#     return t

# def normalize(s):
#     if s is None:
#         return ""
#     s = str(s).lower()
#     s = re.sub(r"\bin\s+(?:tie[\s-]?up|collaboration|association|partnership)\s+with\b.*$", " ", s)
#     s = s.replace(" & ", " and ")
#     s = re.sub(r"[()]", " ", s)
#     s = re.sub(r"[.!?]", " ", s)
#     s = re.sub(r"\bhons\.?\b", "hons", s)
    
#     abbreviations = {
#         r"\bme\b": "mechanical", r"\bce\b": "civil", r"\bcse\b": "computer science",
#         r"\bece\b": "electronics", r"\bee\b": "electrical", r"\beee\b": "electrical electronics",
#         r"\bit\b": "information technology", r"\bai\b": "artificial intelligence",
#         r"\bml\b": "machine learning", r"\bb\.?\s*com\.?\b": "commerce",
#         r"\bbca\b": "computer applications", r"\bmca\b": "computer applications",
#         r"\bmba\b": "business administration", r"\bbba\b": "business administration",
#         r"\bb\.?\s?a\.?\b": "arts", r"\bm\.?\s?a\.?\b": "arts",
#         r"\bll\.?\s*b\.?\b": "law", r"\bll\.?\s*m\.?\b": "law",
#         r"\bbpt\b": "physiotherapy", r"\bbhmct\b": "hotel management catering technology",
#     }
#     for pattern, replacement in abbreviations.items():
#         s = re.sub(pattern, replacement, s)
        
#     s = re.sub(r"[^a-z0-9]+", " ", s)
#     s = re.sub(r"\s+", " ", s).strip()
#     return s

# def similarity(a_norm, b_norm):
#     if not a_norm or not b_norm:
#         return 0.0
#     ta = {_stem(t) for t in a_norm.split() if t not in STOPWORDS}
#     tb = {_stem(t) for t in b_norm.split() if t not in STOPWORDS}
    
#     if not ta or not tb:
#         return 0.0

#     if ta.issubset(tb) or tb.issubset(ta):
#         overlap = len(ta & tb)
#         longer_len = max(len(ta), len(tb))
#         return 0.85 + 0.15 * (overlap / longer_len)

#     a_filtered = " ".join(sorted(ta))
#     b_filtered = " ".join(sorted(tb))
#     seq_ratio = SequenceMatcher(None, a_filtered, b_filtered).ratio()
#     jaccard = len(ta & tb) / len(ta | tb) if (ta | tb) else 0.0
#     score = 0.35 * seq_ratio + 0.65 * jaccard
#     if jaccard == 0.0:
#         score = min(score, 0.2)
#     return score

# def clean_label(raw):
#     s = str(raw).replace("\r\n", "\n")
#     s = re.sub(r"\nNote:-.*", " ", s, flags=re.S)
#     s = re.sub(r"(From \d+.*onwards)", " ", s, flags=re.I)
#     s = s.replace("\n", " / ")
#     s = re.sub(r"[ \t]+", " ", s).strip()
#     s = re.sub(r"\s*/\s*$", " ", s).strip()
#     return s

# _DEGREE_PREFIX = (
#     r"B\.?\s?Tech\.?|M\.?\s?Tech\.?|B\.?\s?Sc\.?|M\.?\s?Sc\.?|BCA|MCA|BBA|MBA|"
#     r"B\.?\s?Com\.?|M\.?\s?Com\.?|B\.?\s?A\.?|M\.?\s?A\.?|BPA|MPA|"
#     r"B\.?\s?Design|M\.?\s?Design|B\s?Design|Diploma|Bachelor|Master|"
#     r"BPT|MPT|B\.?\s?Pharm\.?|M\.?\s?Pharm\.?|B\.?\s?Ed\.?|M\.?\s?Ed\.?|Integrated"
# )
# _DEGREE_PREFIX_RE = re.compile(r"^\s*(?:" + _DEGREE_PREFIX + r")\b", re.I)
# _SEGMENT_SPLIT_RE = re.compile(r"[/,]\s*(?=(?:" + _DEGREE_PREFIX + r"))", re.I)

# def split_segments(clean):
#     parts = [p.strip() for p in _SEGMENT_SPLIT_RE.split(clean) if p.strip()]
#     parts = parts or [clean]
#     expanded = list(parts)
#     for seg in parts:
#         for pm in re.finditer(r"\(([^)]+/[^)]+)\)", seg):
#             inner = pm.group(1)
#             if re.search(r"\bexcept\b", inner, re.I):
#                 continue
#             alts = [a.strip() for a in inner.split("/") if a.strip()]
#             if len(alts) < 2:
#                 continue
#             if all(re.sub(r"\s", "", a).isupper() or len(re.sub(r"\s", "", a)) <= 4 for a in alts):
#                 continue
#             for alt in alts:
#                 expanded.append(seg[:pm.start()] + f"({alt})" + seg[pm.end():])
#     return expanded

# def load_programme_list(xlsx_path):
#     wb = openpyxl.load_workbook(xlsx_path, data_only=True)
#     target_sheet = None
#     header_row_idx = None
#     col = {}
    
#     for sheet_name in wb.sheetnames:
#         ws = wb[sheet_name]
#         for r in range(1, min(ws.max_row, 5) + 1):
#             values = [c.value for c in ws[r]]
#             headers = [str(v).strip().lower() if v else "" for v in values]
#             if any("eligibility" in h for h in headers) and any("programme name" in h or "program name" in h for h in headers):
#                 target_sheet = sheet_name
#                 header_row_idx = r
#                 for i, h in enumerate(headers):
#                     if "eligibility" in h and "elig" not in col:
#                         col["elig"] = i
#                     if ("programme name" in h or "program name" in h) and "name" not in col:
#                         col["name"] = i
#                     if "programme code" in h or "program code" in h or h == "code":
#                         col.setdefault("code", i)
#                 break
#         if target_sheet:
#             break
    
#     if target_sheet is None or "elig" not in col or "name" not in col or "code" not in col:
#         raise RuntimeError("Could not locate the programme-list sheet/columns automatically.")
    
#     ws = wb[target_sheet]
#     programmes = []
#     for r in range(header_row_idx + 1, ws.max_row + 1):
#         elig = ws.cell(r, col["elig"] + 1).value
#         name = ws.cell(r, col["name"] + 1).value
#         code = ws.cell(r, col["code"] + 1).value
#         if code is None or name is None:
#             continue
#         elig = str(elig).strip() if elig else ""
#         name = str(name).strip()
#         code = str(code).strip()
#         if not code or not name:
#             continue
#         if "ph.d" in elig.lower() or "ph.d" in name.lower():
#             continue
#         programmes.append({"EligibilityBase": elig, "ProgrammeName": name, "OfficialCode": code})
    
#     return programmes, wb

# def _leading_token(text_norm):
#     for t in text_norm.split():
#         if t not in STOPWORDS:
#             return t
#     return None

# def _leading_tokens_agree(a, b):
#     if not a or not b:
#         return False
#     if a == b:
#         return True
#     shorter, longer = sorted([a, b], key=len)
#     return len(shorter) >= 3 and longer.startswith(shorter)

# # ---------------------------------------------------------------------------
# # Branch-aware matching helpers
# # ---------------------------------------------------------------------------
# def _extract_branch_keywords(name_norm):
#     keywords = set()
#     for branch, kws in BRANCH_KEYWORDS.items():
#         for kw in kws:
#             if kw in name_norm:
#                 keywords.add(branch)
#                 break
#     return keywords

# def _branch_overlap(keywords1, keywords2):
#     if not keywords1 or not keywords2:
#         return 0.0
#     overlap = len(keywords1 & keywords2)
#     if overlap > 0:
#         return 0.4
#     return 0.0

# def _specificity_score(label_norm):
#     score = 0.0
#     if "(" in label_norm:
#         score += 10
#     specific_kws = ['ai', 'ml', 'data science', 'cyber', 'block chain', 'robotics',
#                     'semiconductor', 'vlsi', 'cloud', 'generative', 'mechatronics',
#                     'aerospace', 'automobile', 'financial markets', 'business analytics',
#                     'international', 'hons', 'research', 'data engineering', 'data analytics']
#     for kw in specific_kws:
#         if kw in label_norm:
#             score += 2
#     score += len(label_norm.split()) * 0.3
#     return score

# def _is_parent_category(parent_label_norm, child_name_norm):
#     for pattern, category in PARENT_PATTERNS:
#         if re.search(pattern, parent_label_norm):
#             if category == 'btech_cse_general':
#                 return any(kw in child_name_norm for kw in ['computer science', 'cse', 'ece', 'electronics', 'electrical', 'information technology', 'robotics'])
#             elif category == 'btech_mech_general':
#                 return any(kw in child_name_norm for kw in ['mechanical', 'civil', 'chemical', 'automobile', 'aerospace', 'mechatronics'])
#             elif category == 'btech_bio_general':
#                 return any(kw in child_name_norm for kw in ['biotechnology', 'biomedical', 'food tech', 'genetics'])
#             elif category == 'mba_general':
#                 return 'mba' in child_name_norm
#             elif category == 'bba_general':
#                 return 'bba' in child_name_norm
#             elif category == 'bcom_general':
#                 return 'commerce' in child_name_norm
#     return False

# # ---------------------------------------------------------------------------
# # Eligibility-workbook-specific loading
# # ---------------------------------------------------------------------------
# CATEGORY_RE = re.compile(r"^[A-Z0-9 .,/()&-]+$")

# _COLUMN_KEYWORDS = {
#     "name": ("programme", "program"),
#     "indian": ("indian",),
#     "international": ("international",),
#     "relaxation": ("relax",),
#     "remarks": ("remark",),
#     "equivalent": ("equivalent",),
# }

# def _detect_eligibility_columns(ws, max_scan_rows=6):
#     for r in range(1, min(ws.max_row, max_scan_rows) + 1):
#         values = [ws.cell(r, c).value for c in range(1, ws.max_column + 1)]
#         headers = [str(v).strip().lower() if v else "" for v in values]
#         col = {}
#         for key, kws in _COLUMN_KEYWORDS.items():
#             for i, h in enumerate(headers):
#                 if any(kw in h for kw in kws):
#                     col.setdefault(key, i)
#                     break
#         if "name" in col and "indian" in col and "international" in col:
#             return r, col
#     return None, {}

# def _sheet_descriptor_text(ws, sheet_name, max_scan_rows=4):
#     bits = [sheet_name]
#     for r in range(1, min(ws.max_row, max_scan_rows) + 1):
#         v = ws.cell(r, 1).value
#         if v:
#             bits.append(str(v))
#     return " ".join(bits).lower()

# def _cell_text(ws, row, col_idx):
#     if col_idx is None:
#         return ""
#     v = ws.cell(row, col_idx + 1).value
#     return str(v).strip() if v is not None and str(v).strip() else ""

# def load_eligibility_workbook(xlsx_path):
#     wb = openpyxl.load_workbook(xlsx_path, data_only=True)
#     sheets = {}
#     for sn in wb.sheetnames:
#         ws = wb[sn]
#         header_row, col = _detect_eligibility_columns(ws)
#         if header_row is None:
#             continue
#         sheets[sn] = {
#             "header_row": header_row,
#             "col": col,
#             "descriptor": _sheet_descriptor_text(ws, sn),
#         }
#     return wb, sheets

# def eligibility_base_to_sheet(elig, sheets):
#     e = elig.lower()
#     def find(keywords):
#         for sn, meta in sheets.items():
#             if any(k in meta["descriptor"] for k in keywords):
#                 return sn
#         return None
#     if "10+2" in e or "12th" in e:
#         return find(["10+2", "12th"])
#     if "10th" in e:
#         return find(["10th"])
#     if "lateral" in e:
#         return find(["lateral"])
#     if "graduation" in e:
#         return find(["graduation"])
#     return None

# def load_eligibility_rows(xlsx_path):
#     wb, sheets = load_eligibility_workbook(xlsx_path)
#     rows_by_sheet = {}
#     for sn, meta in sheets.items():
#         ws = wb[sn]
#         col = meta["col"]
#         anchors = []
#         current_anchor = None
#         category = None
#         for r in range(meta["header_row"] + 1, ws.max_row + 1):
#             name_text = _cell_text(ws, r, col.get("name"))
#             if not name_text:
#                 continue
#             indian_val = _cell_text(ws, r, col.get("indian"))
#             intl_val = _cell_text(ws, r, col.get("international"))
#             is_blank = not indian_val and not intl_val
#             looks_category = bool(CATEGORY_RE.match(name_text)) and len(name_text) < 60
#             looks_like_programme = bool(_DEGREE_PREFIX_RE.match(name_text))
            
#             if is_blank and looks_like_programme and current_anchor is not None:
#                 clean = clean_label(name_text)
#                 for seg in split_segments(clean):
#                     current_anchor["segments"].append({"text": seg, "match_label": clean})
#                 continue
#             if is_blank:
#                 category = name_text
#                 current_anchor = None
#                 continue
                
#             clean = clean_label(name_text)
#             anchor = {
#                 "sheet": sn,
#                 "category": category,
#                 "row": r,
#                 "clean_label": clean,
#                 "eligibility_indian": indian_val,
#                 "eligibility_international": intl_val,
#                 "relaxation": _cell_text(ws, r, col.get("relaxation")),
#                 "remarks": _cell_text(ws, r, col.get("remarks")),
#                 "equivalent_qualification": _cell_text(ws, r, col.get("equivalent")),
#                 "segments": [{"text": seg, "match_label": clean} for seg in split_segments(clean)],
#             }
#             anchors.append(anchor)
#             current_anchor = anchor
#         rows_by_sheet[sn] = anchors
#     return wb, sheets, rows_by_sheet

# # ---------------------------------------------------------------------------
# # Matching with branch-awareness and specificity
# # ---------------------------------------------------------------------------
# def compute_eligibility_mapping(programmes, sheets, rows_by_sheet):
#     out_rows = []
#     for p in programmes:
#         elig, code, name = p["EligibilityBase"], p["OfficialCode"], p["ProgrammeName"]
#         sheet = eligibility_base_to_sheet(elig, sheets)
#         if sheet is None:
#             out_rows.append(_erow(code, name, elig, None, "NOT_FOUND", f'Could not map EligibilityBase "{elig}" to an eligibility sheet.'))
#             continue
            
#         candidates = rows_by_sheet.get(sheet, [])
#         if not candidates:
#             out_rows.append(_erow(code, name, elig, sheet, "NOT_FOUND", f'No eligibility_rows parsed for sheet "{sheet}".'))
#             continue
        
#         name_norm = normalize(name)
#         prog_branches = _extract_branch_keywords(name_norm)
        
#         scored = []
#         for anchor in candidates:
#             best_score, best_label = 0.0, None
#             for seg in anchor["segments"]:
#                 seg_norm = normalize(seg["text"])
#                 base_sim = similarity(name_norm, seg_norm)
                
#                 elig_branches = _extract_branch_keywords(seg_norm)
#                 branch_bonus = _branch_overlap(prog_branches, elig_branches)
#                 specificity = _specificity_score(seg_norm) * 0.01
                
#                 sc = base_sim + branch_bonus + specificity
                
#                 if sc > best_score:
#                     best_score, best_label = sc, seg["match_label"]
            
#             if best_score <= 0.0:
#                 continue
            
#             family_agrees = _leading_tokens_agree(
#                 _leading_token(name_norm), _leading_token(normalize(anchor["clean_label"]))
#             )
            
#             if best_score >= EXACT_THRESHOLD:
#                 scored.append((2, best_score, anchor, best_label))
#             elif best_score >= FUZZY_THRESHOLD and family_agrees:
#                 scored.append((1, best_score, anchor, best_label))
        
#         if not scored:
#             parent_match = None
#             for anchor in candidates:
#                 anchor_norm = normalize(anchor["clean_label"])
#                 if _is_parent_category(anchor_norm, name_norm):
#                     parent_match = anchor
#                     break
            
#             if parent_match:
#                 out_rows.append(_erow(code, name, elig, sheet, "HIGH", "Matched to parent category.", parent_match, parent_match["clean_label"]))
#             else:
#                 out_rows.append(_erow(code, name, elig, sheet, "NOT_FOUND", "No eligibility row matched."))
#             continue
        
#         top_tier = max(s[0] for s in scored)
#         top_group = sorted((s for s in scored if s[0] == top_tier), key=lambda s: s[1], reverse=True)
#         top_score, _, top_anchor, top_label = top_group[0]
#         second_score = top_group[1][1] if len(top_group) > 1 else 0.0
        
#         if top_tier == 2:
#             out_rows.append(_erow(code, name, elig, sheet, "HIGH", "Exact normalized match.", top_anchor, top_label))
#         elif len(top_group) == 1 or (top_score - second_score) >= AMBIGUOUS_MARGIN:
#             out_rows.append(_erow(code, name, elig, sheet, "HIGH", "Closest eligibility row by name similarity with branch awareness.", top_anchor, top_label))
#         else:
#             specificity_scores = [(s[1], _specificity_score(normalize(s[3])), s) for s in top_group]
#             max_spec = max(sp[1] for sp in specificity_scores)
#             most_specific = [sp for sp in specificity_scores if sp[1] == max_spec]
            
#             if len(most_specific) == 1:
#                 _, _, best_match = most_specific[0]
#                 out_rows.append(_erow(code, name, elig, sheet, "HIGH", "Disambiguated by specificity.", best_match[2], best_match[3]))
#             else:
#                 labels = " | ".join(sorted({t[3] for t in top_group if (top_score - t[1]) < AMBIGUOUS_MARGIN}))
#                 out_rows.append(_erow(code, name, elig, sheet, "AMBIGUOUS", "Multiple equally strong eligibility rows matched.", None, labels))
    
#     return out_rows

# def _erow(code, name, elig, sheet, status, reason, anchor=None, label=None):
#     # SIMPLIFIED OUTPUT: Only the core columns needed for the joiner script
#     row_label = label or (anchor["clean_label"] if anchor else "")
#     return {
#         "official_code": code,
#         "eligibility_base": elig,
#         "row_label": row_label,
#         "status": status,
#         "reason": reason
#     }

# # ---------------------------------------------------------------------------
# # Builder entry point
# # ---------------------------------------------------------------------------
# def find_single_file(folder, pattern="*"):
#     files = [f for f in glob.glob(os.path.join(folder, pattern)) if os.path.isfile(f)]
#     if len(files) == 0:
#         raise FileNotFoundError(f"No file found in {folder}")
#     if len(files) > 1:
#         raise ValueError(f"Expected exactly 1 file in {folder}, found {len(files)}: {files}")
#     return files[0]

# def build_mapping(proglist_path=None, eligibility_path=None, out_csv=OUTPUT_CSV):
#     proglist_path = proglist_path or find_single_file(INCOMING_PROGLIST, "*.xls*")
#     eligibility_path = eligibility_path or find_single_file(INCOMING_ELIGIBILITY_EXCEL, "*.xls*")
    
#     with tempfile.TemporaryDirectory() as workdir:
#         prog_xlsx = ensure_xlsx(proglist_path, workdir)
#         elig_xlsx = ensure_xlsx(eligibility_path, workdir)
#         programmes, prog_wb = load_programme_list(prog_xlsx)
#         elig_wb, sheets, rows_by_sheet = load_eligibility_rows(elig_xlsx)
#         mapping = compute_eligibility_mapping(programmes, sheets, rows_by_sheet)
        
#         os.makedirs(os.path.dirname(out_csv), exist_ok=True)
        
#         # Updated fieldnames to match the expected eligibility_row_mapping.csv format
#         fieldnames = ["official_code", "eligibility_base", "row_label", "status", "reason"]
#         with open(out_csv, "w", newline="", encoding="utf-8") as fo:
#             w = csv.DictWriter(fo, fieldnames=fieldnames)
#             w.writeheader()
#             for r in mapping:
#                 w.writerow(r)
        
#         return mapping, out_csv

# def main():
#     mapping, out_csv = build_mapping()
#     codes = [r["official_code"] for r in mapping]
#     dupes = {c for c in codes if codes.count(c) > 1}
#     status_counts = Counter(r["status"] for r in mapping)
    
#     print(f"Wrote {len(mapping)} rows to {out_csv}")
#     print("Status breakdown:", dict(status_counts))
    
#     if dupes:
#         print(f"WARNING: duplicate official_code(s): {sorted(dupes)}")
    
#     needs_review = [r for r in mapping if r["status"] != "HIGH"]
#     if needs_review:
#         print(f"\n{len(needs_review)} row(s) need review:")
#         for r in needs_review:
#             print(f"  - {r['official_code']}: {r['eligibility_base']} -> {r['row_label']} [{r['status']}]")
    
#     if needs_review:
#         sys.exit(1)

# if __name__ == "__main__":
#     main()
# """
# generate_programme_eligibility_mapping.py (FIXED VERSION)
# Outputs a clean mapping CSV where row_labels exactly match the parser's normalization.
# """
import csv
import glob
import os
import re
import sys
import pandas as pd
import tempfile
from collections import Counter
from difflib import SequenceMatcher
import openpyxl

INCOMING_ELIGIBILITY_EXCEL = "lpu_monitor/incoming/eligibility_excel"
INCOMING_PROGLIST = "lpu_monitor/incoming/programme_list"
OUTPUT_CSV = "lpu_monitor/config/eligibility_row_mapping.csv"

EXACT_THRESHOLD = 0.90
FUZZY_THRESHOLD = 0.55
AMBIGUOUS_MARGIN = 0.06

# Branch keyword groups for matching
BRANCH_KEYWORDS = {
    'cse': ['computer science', 'cse', 'ai and ml', 'artificial intelligence', 'machine learning',
            'data science', 'data engineering', 'data analytics', 'cloud computing', 'generative ai',
            'cyber security', 'block chain', 'robotics', 'software product', 'business systems',
            'system design', 'full stack', 'user experience', 'ux', 'ui', 'big data', 'devops',
            'information technology', 'it'],
    'ece': ['electronics and communication', 'ece', 'semiconductor', 'vlsi', 'wireless', 'electronics'],
    'ee': ['electrical engineering', 'ee', 'electrical'],
    'eee': ['electrical and electronics', 'eee'],
    'mech': ['mechanical', 'mechatronics', 'aerospace', 'automobile', 'me'],
    'civil': ['civil', 'construction', 'ce'],
    'chem': ['chemical'],
    'bio': ['biotechnology', 'biomedical', 'food tech', 'genetics', 'bt', 'biotech'],
    'mba': ['mba', 'business administration'],
    'bba': ['bba', 'business administration'],
    'bcom': ['b.com', 'commerce'],
    'bca': ['bca', 'computer applications'],
    'mca': ['mca', 'computer applications'],
    'hotel': ['hotel', 'hospitality', 'catering', 'bhmct', 'tourism'],
    'design': ['design', 'fashion', 'interior', 'animation', 'vfx', 'graphics', 'gaming', 'product', 'industrial'],
    'law': ['law', 'llb', 'llm'],
    'pharma': ['pharmacy', 'pharmaceutics', 'pharmacology', 'bpharm', 'mpharm', 'pharm d'],
}

PARENT_PATTERNS = [
    (r'b\.?\s?tech.*cse.*ece.*ee.*eee.*it', 'btech_cse_general'),
    (r'b\.?\s?tech.*chemical.*civil.*mechanical', 'btech_mech_general'),
    (r'b\.?\s?tech.*biotechnology.*food.*biomedical', 'btech_bio_general'),
    (r'mba.*\(2yrs\)', 'mba_general'),
    (r'mba.*\(2\s?yrs\)', 'mba_general'),
    (r'bba.*\(3yrs\)', 'bba_general'),
    (r'bba.*\(3\s?yrs\)', 'bba_general'),
    (r'b\.?\s?com\.\s*$', 'bcom_general'),
    (r'bachelor.*commerce(?!.*hons)', 'bcom_general'),
]

def ensure_xlsx(path, workdir):
    if path.lower().endswith(".xlsx"):
        return path
    if not path.lower().endswith(".xls"):
        raise ValueError(f"Unsupported file type: {path}")
    os.makedirs(workdir, exist_ok=True)
    out_name = os.path.splitext(os.path.basename(path))[0] + ".xlsx"
    out_path = os.path.join(workdir, out_name)
    xls = pd.ExcelFile(path, engine="xlrd")
    with pd.ExcelWriter(out_path, engine="openpyxl") as writer:
        for sheet_name in xls.sheet_names:
            df = xls.parse(sheet_name, header=None)
            df.to_excel(writer, sheet_name=sheet_name, index=False, header=False)
    return out_path

STOPWORDS = {
    "b", "m", "of", "and", "the", "in", "for", "a", "an", "with", "hons",
    "bachelor", "bachelors", "master", "masters", "programme", "program",
    "technology", "engineering", "science", "sciences", "studies", "design",
    "no", "integrated", "integr", "including", "option", "transfer", "credit",
    "after", "exit", "discontinue", "lateral", "entry"
}

def _stem(t):
    if len(t) > 4 and t.endswith("s") and not t.endswith("ss"):
        return t[:-1]
    return t

def normalize(s):
    """
    CRITICAL FIX: Removed abbreviation expansions that altered the text (e.g., mba -> business administration).
    Uses the EXACT SAME normalization logic as the parser to guarantee CSV matches.
    """
    if s is None:
        return ""
    s = str(s)
    s = re.split(r"note\s*:", s, flags=re.IGNORECASE)[0]
    s = s.replace("\n", " ").replace("\r", " ")
    s = re.sub(r"[^a-z0-9]+", " ", s.lower())
    s = re.sub(r"\s+", " ", s).strip()
    return s

def similarity(a_norm, b_norm):
    if not a_norm or not b_norm:
        return 0.0
    ta = {_stem(t) for t in a_norm.split() if t not in STOPWORDS}
    tb = {_stem(t) for t in b_norm.split() if t not in STOPWORDS}
    if not ta or not tb:
        return 0.0
    if ta.issubset(tb) or tb.issubset(ta):
        overlap = len(ta & tb)
        longer_len = max(len(ta), len(tb))
        return 0.85 + 0.15 * (overlap / longer_len)
    a_filtered = " ".join(sorted(ta))
    b_filtered = " ".join(sorted(tb))
    seq_ratio = SequenceMatcher(None, a_filtered, b_filtered).ratio()
    jaccard = len(ta & tb) / len(ta | tb) if (ta | tb) else 0.0
    score = 0.35 * seq_ratio + 0.65 * jaccard
    if jaccard == 0.0:
        score = min(score, 0.2)
    return score

def clean_label(raw):
    s = str(raw).replace("\r\n", "\n")
    s = re.sub(r"\nNote:-.*", " ", s, flags=re.S)
    s = re.sub(r"(From \d+.*onwards)", " ", s, flags=re.I)
    s = s.replace("\n", " / ")
    s = re.sub(r"[ \t]+", " ", s).strip()
    s = re.sub(r"\s*/\s*$", " ", s).strip()
    return s

_DEGREE_PREFIX = (
    r"B\.?\s?Tech\.?|M\.?\s?Tech\.?|B\.?\s?Sc\.?|M\.?\s?Sc\.?|BCA|MCA|BBA|MBA|"
    r"B\.?\s?Com\.?|M\.?\s?Com\.?|B\.?\s?A\.?|M\.?\s?A\.?|BPA|MPA|"
    r"B\.?\s?Design|M\.?\s?Design|B\s?Design|Diploma|Bachelor|Master|"
    r"BPT|MPT|B\.?\s?Pharm\.?|M\.?\s?Pharm\.?|B\.?\s?Ed\.?|M\.?\s?Ed\.?|Integrated"
)
_DEGREE_PREFIX_RE = re.compile(r"^\s*(?:" + _DEGREE_PREFIX + r")\b", re.I)
_SEGMENT_SPLIT_RE = re.compile(r"[/,]\s*(?=(?:" + _DEGREE_PREFIX + r"))", re.I)

def split_segments(clean):
    parts = [p.strip() for p in _SEGMENT_SPLIT_RE.split(clean) if p.strip()]
    parts = parts or [clean]
    expanded = list(parts)
    for seg in parts:
        for pm in re.finditer(r"\(([^)]+/[^)]+)\)", seg):
            inner = pm.group(1)
            if re.search(r"\bexcept\b", inner, re.I):
                continue
            alts = [a.strip() for a in inner.split("/") if a.strip()]
            if len(alts) < 2:
                continue
            if all(re.sub(r"\s", "", a).isupper() or len(re.sub(r"\s", "", a)) <= 4 for a in alts):
                continue
            for alt in alts:
                expanded.append(seg[:pm.start()] + f"({alt})" + seg[pm.end():])
    return expanded

def load_programme_list(xlsx_path):
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    target_sheet = None
    header_row_idx = None
    col = {}
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        for r in range(1, min(ws.max_row, 5) + 1):
            values = [c.value for c in ws[r]]
            headers = [str(v).strip().lower() if v else "" for v in values]
            if any("eligibility" in h for h in headers) and any("programme name" in h or "program name" in h for h in headers):
                target_sheet = sheet_name
                header_row_idx = r
                for i, h in enumerate(headers):
                    if "eligibility" in h and "elig" not in col: col["elig"] = i
                    if ("programme name" in h or "program name" in h) and "name" not in col: col["name"] = i
                    if "programme code" in h or "program code" in h or h == "code": col.setdefault("code", i)
                break
        if target_sheet: break
    if target_sheet is None or "elig" not in col or "name" not in col or "code" not in col:
        raise RuntimeError("Could not locate the programme-list sheet/columns automatically.")
    ws = wb[target_sheet]
    programmes = []
    for r in range(header_row_idx + 1, ws.max_row + 1):
        elig = ws.cell(r, col["elig"] + 1).value
        name = ws.cell(r, col["name"] + 1).value
        code = ws.cell(r, col["code"] + 1).value
        if code is None or name is None: continue
        elig = str(elig).strip() if elig else ""
        name = str(name).strip()
        code = str(code).strip()
        if not code or not name: continue
        if "ph.d" in elig.lower() or "ph.d" in name.lower(): continue
        programmes.append({"EligibilityBase": elig, "ProgrammeName": name, "OfficialCode": code})
    return programmes, wb

def _leading_token(text_norm):
    for t in text_norm.split():
        if t not in STOPWORDS: return t
    return None

def _leading_tokens_agree(a, b):
    if not a or not b: return False
    if a == b: return True
    shorter, longer = sorted([a, b], key=len)
    return len(shorter) >= 3 and longer.startswith(shorter)
def _expand_it_synonym(text_norm):
    """Ensure 'it' and 'information' co-occur so BSc(IT) matches B.Sc.(Information Technology)."""
    has_it = bool(re.search(r'\bit\b', text_norm))
    has_info = bool(re.search(r'\binformation\b', text_norm))
    if has_it and not has_info:
        return text_norm + " information"
    if has_info and not has_it:
        return text_norm + " it"
    return text_norm

def _extract_branch_keywords(name_norm):
    keywords = set()
    for branch, kws in BRANCH_KEYWORDS.items():
        for kw in kws:
            if kw in name_norm:
                keywords.add(branch)
                break
    return keywords

def _branch_overlap(keywords1, keywords2):
    if not keywords1 or not keywords2: return 0.0
    overlap = len(keywords1 & keywords2)
    if overlap > 0: return 0.4
    return 0.0

def _specificity_score(label_norm):
    score = 0.0
    if "(" in label_norm: score += 10
    specific_kws = ['ai', 'ml', 'data science', 'cyber', 'block chain', 'robotics',
                    'semiconductor', 'vlsi', 'cloud', 'generative', 'mechatronics',
                    'aerospace', 'automobile', 'financial markets', 'business analytics',
                    'international', 'hons', 'research', 'data engineering', 'data analytics']
    for kw in specific_kws:
        if kw in label_norm: score += 2
    score += len(label_norm.split()) * 0.3
    return score
def _extract_acronym_from_raw(raw_text):
    """Extract short acronym from parentheses in raw (un-normalized) text.
    E.g. 'Bachelor of Hotel Mgt (BHMCT)' -> 'bhmct'
         'MBA (2yrs)'                     -> None  (starts with digit)
         'B.Tech (CSE - AI and ML)'       -> None  (too long/complex)
    """
    lowered = raw_text.lower()
    for m in re.finditer(r'\(([a-z][a-z0-9.]*|[a-z]+ [a-z0-9.]+)\)', lowered):
        acronym = m.group(1).strip().replace(".", "")
        words = acronym.split()
        if len(words) <= 2 and len(acronym) <= 12:
            return acronym
    return None

def _is_parent_category(parent_label_norm, child_name_norm):
    for pattern, category in PARENT_PATTERNS:
        if re.search(pattern, parent_label_norm):
            if category == 'btech_cse_general':
                return any(kw in child_name_norm for kw in ['computer science', 'cse', 'ece', 'electronics', 'electrical', 'information technology', 'robotics'])
            elif category == 'btech_mech_general':
                return any(kw in child_name_norm for kw in ['mechanical', 'civil', 'chemical', 'automobile', 'aerospace', 'mechatronics'])
            elif category == 'btech_bio_general':
                return any(kw in child_name_norm for kw in ['biotechnology', 'biomedical', 'food tech', 'genetics'])
            elif category == 'mba_general': return 'mba' in child_name_norm
            elif category == 'bba_general': return 'bba' in child_name_norm
            elif category == 'bcom_general': return 'commerce' in child_name_norm
    return False

CATEGORY_RE = re.compile(r"^[A-Z0-9 .,/()&-]+$")
_COLUMN_KEYWORDS = {
    "name": ("programme", "program"), "indian": ("indian",), "international": ("international",),
    "relaxation": ("relax",), "remarks": ("remark",), "equivalent": ("equivalent",),
}

def _detect_eligibility_columns(ws, max_scan_rows=6):
    for r in range(1, min(ws.max_row, max_scan_rows) + 1):
        values = [ws.cell(r, c).value for c in range(1, ws.max_column + 1)]
        headers = [str(v).strip().lower() if v else "" for v in values]
        col = {}
        for key, kws in _COLUMN_KEYWORDS.items():
            for i, h in enumerate(headers):
                if any(kw in h for kw in kws): col.setdefault(key, i); break
        if "name" in col and "indian" in col and "international" in col: return r, col
    return None, {}

def _sheet_descriptor_text(ws, sheet_name, max_scan_rows=4):
    bits = [sheet_name]
    for r in range(1, min(ws.max_row, max_scan_rows) + 1):
        v = ws.cell(r, 1).value
        if v: bits.append(str(v))
    return " ".join(bits).lower()

def _cell_text(ws, row, col_idx):
    if col_idx is None: return ""
    v = ws.cell(row, col_idx + 1).value
    return str(v).strip() if v is not None and str(v).strip() else ""

def load_eligibility_workbook(xlsx_path):
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    sheets = {}
    for sn in wb.sheetnames:
        ws = wb[sn]
        header_row, col = _detect_eligibility_columns(ws)
        if header_row is None: continue
        sheets[sn] = {"header_row": header_row, "col": col, "descriptor": _sheet_descriptor_text(ws, sn)}
    return wb, sheets

def eligibility_base_to_sheet(elig, sheets):
    e = elig.lower()
    def find(keywords):
        for sn, meta in sheets.items():
            if any(k in meta["descriptor"] for k in keywords): return sn
        return None
    if "10+2" in e or "12th" in e: return find(["10+2", "12th"])
    if "10th" in e: return find(["10th"])
    if "lateral" in e: return find(["lateral"])
    if "graduation" in e: return find(["graduation"])
    return None

def load_eligibility_rows(xlsx_path):
    wb, sheets = load_eligibility_workbook(xlsx_path)
    rows_by_sheet = {}
    for sn, meta in sheets.items():
        ws = wb[sn]
        col = meta["col"]
        anchors = []
        current_anchor = None
        category = None
        for r in range(meta["header_row"] + 1, ws.max_row + 1):
            name_text = _cell_text(ws, r, col.get("name"))
            if not name_text: continue
            indian_val = _cell_text(ws, r, col.get("indian"))
            intl_val = _cell_text(ws, r, col.get("international"))
            is_blank = not indian_val and not intl_val
            looks_like_programme = bool(_DEGREE_PREFIX_RE.match(name_text))
            if is_blank and looks_like_programme and current_anchor is not None:
                clean = clean_label(name_text)
                for seg in split_segments(clean):
                    current_anchor["segments"].append({"text": seg, "match_label": clean})
                continue
            if is_blank:
                category = name_text
                current_anchor = None
                continue
            clean = clean_label(name_text)
            anchor = {
                "sheet": sn, "category": category, "row": r, "clean_label": clean,
                "eligibility_indian": indian_val, "eligibility_international": intl_val,
                "relaxation": _cell_text(ws, r, col.get("relaxation")),
                "remarks": _cell_text(ws, r, col.get("remarks")),
                "equivalent_qualification": _cell_text(ws, r, col.get("equivalent")),
                "segments": [{"text": seg, "match_label": clean} for seg in split_segments(clean)],
            }
            anchors.append(anchor)
            current_anchor = anchor
        rows_by_sheet[sn] = anchors
    return wb, sheets, rows_by_sheet

def compute_eligibility_mapping(programmes, sheets, rows_by_sheet):
    out_rows = []
    for p in programmes:
        elig, code, name = p["EligibilityBase"], p["OfficialCode"], p["ProgrammeName"]
        sheet = eligibility_base_to_sheet(elig, sheets)
        if sheet is None:
            out_rows.append(_erow(code, name, elig, None, "NOT_FOUND", f'Could not map EligibilityBase "{elig}" to an eligibility sheet.'))
            continue
        candidates = rows_by_sheet.get(sheet, [])
        if not candidates:
            out_rows.append(_erow(code, name, elig, sheet, "NOT_FOUND", f'No eligibility_rows parsed for sheet "{sheet}".'))
            continue
        
        name_norm = normalize(name)
        name_acronym = _extract_acronym_from_raw(name)   # <-- ADD THIS LINE
        prog_branches = _extract_branch_keywords(name_norm)
        
        scored = []
        for anchor in candidates:
            best_score, best_label = 0.0, None
            for seg in anchor["segments"]:
                seg_norm = normalize(seg["text"])
                base_sim = similarity(_expand_it_synonym(name_norm), _expand_it_synonym(seg_norm))
                elig_branches = _extract_branch_keywords(seg_norm)
                branch_bonus = _branch_overlap(prog_branches, elig_branches)
                specificity = _specificity_score(seg_norm) * 0.01
                sc = base_sim + branch_bonus + specificity
                # <-- ADD THESE 3 LINES
                if name_acronym and (seg_norm.startswith(name_acronym + " ") or seg_norm == name_acronym):
                    sc += 0.5
                if sc > best_score:
                    best_score, best_label = sc, seg["match_label"]
            if best_score <= 0.0: continue
            family_agrees = _leading_tokens_agree(_leading_token(name_norm), _leading_token(normalize(anchor["clean_label"])))
            if best_score >= EXACT_THRESHOLD:
                scored.append((2, best_score, anchor, best_label))
            elif best_score >= FUZZY_THRESHOLD and family_agrees:
                scored.append((1, best_score, anchor, best_label))
        
        if not scored:
            parent_match = None
            for anchor in candidates:
                anchor_norm = normalize(anchor["clean_label"])
                if _is_parent_category(anchor_norm, name_norm):
                    parent_match = anchor
                    break
            if parent_match:
                out_rows.append(_erow(code, name, elig, sheet, "HIGH", "Matched to parent category.", parent_match, parent_match["clean_label"]))
            else:
                out_rows.append(_erow(code, name, elig, sheet, "NOT_FOUND", "No eligibility row matched."))
            continue
        
        top_tier = max(s[0] for s in scored)
        top_group = sorted((s for s in scored if s[0] == top_tier), key=lambda s: s[1], reverse=True)
        top_score, _, top_anchor, top_label = top_group[0]
        second_score = top_group[1][1] if len(top_group) > 1 else 0.0
        
        if top_tier == 2:
            out_rows.append(_erow(code, name, elig, sheet, "HIGH", "Exact normalized match.", top_anchor, top_label))
        elif len(top_group) == 1 or (top_score - second_score) >= AMBIGUOUS_MARGIN:
            out_rows.append(_erow(code, name, elig, sheet, "HIGH", "Closest eligibility row by name similarity with branch awareness.", top_anchor, top_label))
        else:
            specificity_scores = [(s[1], _specificity_score(normalize(s[3])), s) for s in top_group]
            max_spec = max(sp[1] for sp in specificity_scores)
            most_specific = [sp for sp in specificity_scores if sp[1] == max_spec]
            if len(most_specific) == 1:
                _, _, best_match = most_specific[0]
                out_rows.append(_erow(code, name, elig, sheet, "HIGH", "Disambiguated by specificity.", best_match[2], best_match[3]))
            else:
                labels = " | ".join(sorted({t[3] for t in top_group if (top_score - t[1]) < AMBIGUOUS_MARGIN}))
                out_rows.append(_erow(code, name, elig, sheet, "AMBIGUOUS", "Multiple equally strong eligibility rows matched.", None, labels))
    return out_rows

def _erow(code, name, elig, sheet, status, reason, anchor=None, label=None):
    row_label = label or (anchor["clean_label"] if anchor else "")
    # CRITICAL FIX: Apply the exact same normalization to the output CSV as the parser uses.
    return {
        "official_code": code,
        "eligibility_base": elig,
        "row_label": normalize(row_label), 
        "status": status,
        "reason": reason
    }

def find_single_file(folder, pattern="*"):
    files = [f for f in glob.glob(os.path.join(folder, pattern)) if os.path.isfile(f)]
    if len(files) == 0: raise FileNotFoundError(f"No file found in {folder}")
    if len(files) > 1: raise ValueError(f"Expected exactly 1 file in {folder}, found {len(files)}: {files}")
    return files[0]

def build_mapping(proglist_path=None, eligibility_path=None, out_csv=OUTPUT_CSV):
    proglist_path = proglist_path or find_single_file(INCOMING_PROGLIST, "*.xls*")
    eligibility_path = eligibility_path or find_single_file(INCOMING_ELIGIBILITY_EXCEL, "*.xls*")
    with tempfile.TemporaryDirectory() as workdir:
        prog_xlsx = ensure_xlsx(proglist_path, workdir)
        elig_xlsx = ensure_xlsx(eligibility_path, workdir)
        programmes, prog_wb = load_programme_list(prog_xlsx)
        elig_wb, sheets, rows_by_sheet = load_eligibility_rows(elig_xlsx)
        mapping = compute_eligibility_mapping(programmes, sheets, rows_by_sheet)
        os.makedirs(os.path.dirname(out_csv), exist_ok=True)
        fieldnames = ["official_code", "eligibility_base", "row_label", "status", "reason"]
        with open(out_csv, "w", newline="", encoding="utf-8") as fo:
            w = csv.DictWriter(fo, fieldnames=fieldnames)
            w.writeheader()
            for r in mapping: w.writerow(r)
        return mapping, out_csv

def main():
    mapping, out_csv = build_mapping()
    codes = [r["official_code"] for r in mapping]
    dupes = {c for c in codes if codes.count(c) > 1}
    status_counts = Counter(r["status"] for r in mapping)
    print(f"Wrote {len(mapping)} rows to {out_csv}")
    print("Status breakdown:", dict(status_counts))
    if dupes: print(f"WARNING: duplicate official_code(s): {sorted(dupes)}")
    needs_review = [r for r in mapping if r["status"] != "HIGH"]
    if needs_review:
        print(f"\n{len(needs_review)} row(s) need review:")
        for r in needs_review: print(f"  - {r['official_code']}: {r['eligibility_base']} -> {r['row_label']} [{r['status']}]")
    if needs_review: sys.exit(1)

if __name__ == "__main__":
    main()
