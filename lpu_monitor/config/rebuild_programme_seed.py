"""
rebuild_programme_seed.py

Regenerates programme_seed.csv from the ProgramId-source file (the one with
OfficialCode + ProgramId + ProgramName columns -- a different file from the
Final Programme List Excel, which has Annexure/Discipline data but NOT
ProgramId). Programmes change year to year, so this needs to be rerunnable
any time a fresh source file arrives, not a one-time build.
"""
import openpyxl
import csv
import glob
import os


def find_single_file(folder, pattern="*.xlsx"):
    files = glob.glob(os.path.join(folder, pattern))
    if len(files) == 0:
        raise FileNotFoundError(f"No file found in {folder}")
    if len(files) > 1:
        raise ValueError(f"Expected exactly 1 file in {folder}, found {len(files)}: {files}")
    return files[0]


def rebuild_programme_seed(xlsx_path, out_csv="lpu_monitor/config/programme_seed.csv"):
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    ws = wb.active
    header = [c.value for c in ws[1]]

    required = ["OfficialCode", "ProgramId", "ProgramName"]
    missing = [c for c in required if c not in header]
    if missing:
        raise ValueError(
            f"Expected columns {required} not found. Missing: {missing}. "
            f"Actual columns present: {header} -- file structure may have changed."
        )

    idx = {c: header.index(c) for c in required}
    rows = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        code = row[idx["OfficialCode"]]
        if not code:
            continue
        rows.append({
            "OfficialCode": code,
            "ProgramId": row[idx["ProgramId"]],
            "ProgramName": row[idx["ProgramName"]],
        })

    with open(out_csv, "w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=["OfficialCode", "ProgramId", "ProgramName"])
        writer.writeheader()
        writer.writerows(rows)

    return len(rows)


if __name__ == "__main__":
    import sys
    if len(sys.argv) > 1:
        xlsx_path = sys.argv[1]
    else:
        xlsx_path = find_single_file("lpu_monitor/incoming/program_id/")

    count = rebuild_programme_seed(xlsx_path)
    print(f"Rebuilt programme_seed.csv with {count} programmes.")