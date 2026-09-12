"""
lpu_monitor/config/diagnose_p344.py
"""
import glob
import os
import openpyxl

# Find programme list
folder = "lpu_monitor/incoming/programme_list"
files = [f for f in glob.glob(os.path.join(folder, "*.xls*")) if os.path.isfile(f)]
if not files:
    print("No programme list file found!")
    exit(1)

# Convert to xlsx if needed
import tempfile, pandas as pd
with tempfile.TemporaryDirectory() as td:
    p = files[0]
    if p.lower().endswith(".xls") and not p.lower().endswith(".xlsx"):
        xls = pd.ExcelFile(p, engine="xlrd")
        p2 = os.path.join(td, "prog.xlsx")
        with pd.ExcelWriter(p2, engine="openpyxl") as w:
            for sn in xls.sheet_names:
                xls.parse(sn, header=None).to_excel(w, sheet_name=sn, index=False, header=False)
        p = p2

    wb = openpyxl.load_workbook(p, data_only=True)
    for sn in wb.sheetnames:
        ws = wb[sn]
        for r in range(1, ws.max_row + 1):
            # Find the code column
            for c in range(1, ws.max_column + 1):
                v = ws.cell(r, c).value
                if v and str(v).strip() == "P344":
                    # Print the entire row
                    headers = [ws.cell(1, cc).value for cc in range(1, ws.max_column + 1)]
                    vals = [ws.cell(r, cc).value for cc in range(1, ws.max_column + 1)]
                    print(f"Found P344 in sheet '{sn}', row {r}:")
                    for h, v in zip(headers, vals):
                        if h and str(h).strip():
                            print(f"  {str(h).strip()}: {v}")
                    print()
                    exit(0)
    print("P344 not found in any sheet!")