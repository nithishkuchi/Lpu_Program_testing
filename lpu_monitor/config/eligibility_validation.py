"""
eligibility_validation.py

Validates eligibility text (percentage cutoffs, required subjects, accepted
entrance exams) between the expected Eligibility Excel and the live page's
"eligibilityindian" field.

WHY NOT FUZZY / SEMANTIC MATCHING:
The expected text ("60% in 10+2 with (English, Physics & Mathematics) +
LPUNEST/ JEE Main/ CUET") and the live text (a much longer HTML paragraph
full of popup/tooltip text, disclaimers, and formatting) will almost never
share matching sentence structure -- that part of the earlier concern is
correct. But the deeper problem with word-fuzzy OR whole-paragraph semantic
similarity here isn't just "words don't match" -- it's that a SINGLE
important fact (a percentage cutoff, a required subject) is often a tiny
fraction of a long paragraph. A semantic model comparing the whole
paragraph would likely still score two paragraphs as "very similar" even
if that one fact changed (e.g. 60% -> 65%), because the other 95% of the
paragraph (boilerplate, tooltips, disclaimers) stayed identical. That's the
opposite failure mode from Important Dates' title-matching problem: there
the risk was two DIFFERENT things scoring falsely similar; here the risk is
a REAL change getting diluted into invisibility by everything around it.

INSTEAD: extract the specific, objectively-checkable facts that actually
carry meaning here -- percentage cutoffs, required subjects, accepted
entrance exams -- via simple regex/keyword extraction, and compare THOSE
extracted fact-sets directly. This sidesteps paragraph-level wording
differences entirely, because nothing here is comparing sentences.

ON THE SUBJECT/EXAM VOCABULARY LISTS BELOW: these are NOT the kind of
per-programme dynamic content (dates, programme names) that should never be
hardcoded -- they're a small, slow-changing REFERENCE vocabulary (there are
only so many entrance exams and school subjects LPU's eligibility criteria
ever refer to). If a new exam or subject name shows up that isn't in these
lists, it just won't be extracted/compared -- add it to the relevant set
below, same as GENERIC_TITLE_WORDS was maintained for Important Dates.

CATEGORY ALIASES (e.g. "Basic Science" meaning Physics/Chemistry/Maths) ARE
dynamic, sourced from the "Medical Terminology PG Sciences" sheet each
year -- see load_category_aliases() below, which reads that sheet directly
rather than hardcoding its contents.
"""
import re


def strip_html(text):
    if text is None:
        return ""
    return re.sub(r"<[^>]+>", "", str(text))


def clean_for_display(text):
    text = strip_html(text)
    text = text.replace("\n", " ").replace("\r", " ")
    return re.sub(r"\s+", " ", text).strip()


# ---------------------------------------------------------------------------
# Percentage extraction -- fully generic, no vocabulary needed.
# ---------------------------------------------------------------------------
def extract_percentages(text):
    text = strip_html(text)
    return {float(m) for m in re.findall(r"(\d+(?:\.\d+)?)\s*%", text)}


# ---------------------------------------------------------------------------
# Subject extraction. Base vocabulary covers general school/college subjects
# seen directly in Diploma/After 10+2/After Graduation sheets. Each entry is
# canonical_name -> surface-text variants to search for (case-insensitive,
# whole-word).
# ---------------------------------------------------------------------------
SUBJECT_VARIANTS = {
    "english": ["english"],
    "physics": ["physics"],
    "chemistry": ["chemistry"],
    "mathematics": ["mathematics", "maths", "math"],
    "biology": ["biology", "bio"],
    "biotechnology": ["biotechnology", "biotech"],
    "science": ["science"],
    "computer science": ["computer science"],
    "commerce": ["commerce"],
    "economics": ["economics"],
    "accountancy": ["accountancy"],
    "business studies": ["business studies"],
    "statistics": ["statistics"],
}


def load_category_aliases(xlsx_path, sheet_name="Medical Termiology PG Sciences"):
    """
    Reads the category -> subject-list mapping directly from the sheet each
    year, e.g. "Basic Science" -> {physics, chemistry, mathematics}. This is
    dynamic/sourced data, unlike SUBJECT_VARIANTS above -- these category
    LABELS and their expansions can genuinely change year to year, so they
    are never hardcoded, only read from the file.

    Returns: {canonical_category_label: set(canonical_subject_names)}
    Subject names in the expansion are matched against SUBJECT_VARIANTS
    where possible; anything not in that vocabulary is kept as its own
    lowercased canonical entry (e.g. "botany", "mbbs") so it can still be
    matched literally even though it's not a "general" subject.
    """
    import openpyxl

    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    if sheet_name not in wb.sheetnames:
        return {}
    ws = wb[sheet_name]

    aliases = {}
    for r in range(2, ws.max_row + 1):  # row 1 is a title, not a category
        category = ws.cell(row=r, column=1).value
        subjects_cell = ws.cell(row=r, column=2).value
        if not category or not subjects_cell:
            continue
        category_key = str(category).strip().lower()
        subject_names = {s.strip().lower() for s in str(subjects_cell).split(",") if s.strip()}
        aliases[category_key] = subject_names
    return aliases


def _boundary_pattern(variant):
    """\b only makes sense at a word-char/non-word-char transition -- it
    silently fails to match when a variant's edge is punctuation (e.g. the
    trailing ')' in "jee (main)"), because both sides of that position end
    up non-word. Only wrap \b around edges that are actually alphanumeric;
    let punctuation edges match literally instead."""
    escaped = re.escape(variant)
    prefix = r"\b" if variant[0].isalnum() else ""
    suffix = r"\b" if variant[-1].isalnum() else ""
    return prefix + escaped + suffix


def extract_subjects(text, category_aliases=None):
    """
    Finds every known subject (or subject-category alias, expanded via
    category_aliases) mentioned in the text. Whole-word matching so e.g.
    "math" doesn't match inside an unrelated word.

    Category aliases are checked FIRST, and the matched phrase is masked
    out of the text before generic subject-word matching runs. Without
    this, a generic word that's a literal substring of a category label
    (e.g. "science" inside "Basic Science") gets falsely detected as its
    own separate, unreciprocated subject entry on top of the alias
    expansion -- confirmed as a real false-positive on a test case where
    "Basic Science" (expected) and "Physics, Chemistry and Maths" (live)
    genuinely mean the same thing.
    """
    text_lower = strip_html(text).lower()
    found = set()
    masked = text_lower

    if category_aliases:
        for category_key, subject_set in category_aliases.items():
            match = re.search(_boundary_pattern(category_key), masked)
            if match:
                for subj in subject_set:
                    # Map through SUBJECT_VARIANTS if it's a known general
                    # subject; otherwise keep it as its own literal entry
                    # (e.g. "botany", "mbbs" aren't in SUBJECT_VARIANTS).
                    canonical = next(
                        (c for c, variants in SUBJECT_VARIANTS.items() if subj in variants or subj == c),
                        subj,
                    )
                    found.add(canonical)
                masked = masked[:match.start()] + " " * (match.end() - match.start()) + masked[match.end():]

    for canonical, variants in SUBJECT_VARIANTS.items():
        for variant in variants:
            if re.search(_boundary_pattern(variant), masked):
                found.add(canonical)
                break

    return found


# ---------------------------------------------------------------------------
# Entrance exam extraction -- same maintained-vocabulary approach.
# ---------------------------------------------------------------------------
EXAM_VARIANTS = {
    "lpunest": ["lpunest"],
    "jee main": ["jee main", "jee (main)", "jee-main"],
    "jee advanced": ["jee advanced", "jee (advanced)"],
    "cuet": ["cuet"],
    "neet": ["neet"],
    "cat": ["cat"],
    "mat": ["mat"],
    "xat": ["xat"],
    "cmat": ["cmat"],
    "clat": ["clat"],
    "ielts": ["ielts"],
    "toefl": ["toefl"],
}


def extract_exams(text):
    text_lower = strip_html(text).lower()
    found = set()
    for canonical, variants in EXAM_VARIANTS.items():
        for variant in variants:
            if re.search(_boundary_pattern(variant), text_lower):
                found.add(canonical)
                break
    return found


# ---------------------------------------------------------------------------
# Fact extraction + comparison
# ---------------------------------------------------------------------------
def extract_facts(text, category_aliases=None):
    return {
        "percentages": extract_percentages(text),
        "subjects": extract_subjects(text, category_aliases),
        "exams": extract_exams(text),
    }


FACT_LABELS = {
    "percentages": "Percentage cutoffs",
    "subjects": "Subjects",
    "exams": "Accepted exams",
}


def format_percent_set(values):
    return ", ".join(f"{v:g}%" for v in sorted(values))


def format_name_set(values):
    return ", ".join(sorted(values)) if values else "(none found)"

def _strip_relaxation_sentences(text):
    """Remove sentences containing 'relax' so their percentages aren't
    mistaken for eligibility cutoffs."""
    sentences = re.split(r'(?<=[.!?])\s+|\n', text)
    return " ".join(s for s in sentences if not re.search(r'\brelax\w*\b', s, re.I))
def _strip_scholarship_sentences(text):
    """Remove sentences about scholarships so their exam mentions
    (e.g. 'CUET for Scholarship Details') aren't flagged as eligibility exams."""
    sentences = re.split(r'(?<=[.!?])\s+|\n', text)
    return " ".join(s for s in sentences if not re.search(r'\bscholarship\w*\b', s, re.I))
def compare_eligibility(expected_indian, expected_relaxation, live_eligibility_html,
                        category_aliases=None, expected_equivalent=""):
    expected_text = " ".join(t for t in (expected_indian, expected_relaxation, expected_equivalent) if t)
    live_text = strip_html(live_eligibility_html or "")

    # Percentages: strip relaxation sentences only
    expected_pct_text = _strip_relaxation_sentences(expected_indian or "")
    live_pct_text = _strip_relaxation_sentences(live_text)

    # Exams: strip BOTH relaxation AND scholarship sentences
    expected_exam_text = _strip_relaxation_sentences(_strip_scholarship_sentences(expected_text))
    live_exam_text = _strip_relaxation_sentences(_strip_scholarship_sentences(live_text))

    # Subjects: use full text (subjects aren't in relaxation/scholarship contexts)
    expected_subj_text = expected_text
    live_subj_text = live_text

    expected_facts = {
        "percentages": extract_percentages(expected_pct_text),
        "subjects": extract_subjects(expected_subj_text, category_aliases),
        "exams": extract_exams(expected_exam_text),
    }
    live_facts = {
        "percentages": extract_percentages(live_pct_text),
        "subjects": extract_subjects(live_subj_text, category_aliases),
        "exams": extract_exams(live_exam_text),
    }

    blocks = []
    for key in ("percentages", "subjects", "exams"):
        expected_only = expected_facts[key] - live_facts[key]
        live_only = live_facts[key] - expected_facts[key]
        if not expected_only and not live_only:
            continue

        formatter = format_percent_set if key == "percentages" else format_name_set
        lines = [f"{FACT_LABELS[key]}:"]
        if expected_only:
            lines.append(f"expected: {formatter(expected_only)}")
        if live_only:
            lines.append(f"live: {formatter(live_only)}")
        blocks.append("\n".join(lines))

    return "\n\n".join(blocks)